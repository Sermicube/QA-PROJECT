"""Genera la plantilla sintética CO-FR-VRA-03 para desarrollo y tests.

Ejecutar dentro del contenedor: python scripts/generate_template.py
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def make_template(output_path: Path) -> None:
    wb = openpyxl.Workbook()

    # ── Hoja Pruebas ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Pruebas"

    header_font = Font(bold=True, size=11)
    label_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "PRUEBAS DE CALIDAD DE SOFTWARE"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "CO-FR-VRA-03"

    labels = {
        "A5": "Código de requerimiento:", "C5": "",
        "A6": "Módulo:", "C6": "",
        "A7": "Analista:", "C7": "",
        "A8": "Fecha:", "C8": "",
    }
    for cell, val in labels.items():
        ws[cell] = val
        if cell.startswith("A"):
            ws[cell].font = header_font
            ws[cell].fill = label_fill

    case_headers = ["Código", "Nombre del caso", "Pasos", "Resultado esperado", "Resultado"]
    for col, header in enumerate(case_headers, start=1):
        cell = ws.cell(row=11, column=col, value=header)
        cell.font = header_font
        cell.fill = label_fill
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, horizontal="center")

    for row in range(12, 17):
        for col in range(1, 6):
            ws.cell(row=row, column=col).border = border

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 35
    ws.column_dimensions["E"].width = 12

    # ── Hoja Evidencias ───────────────────────────────────────────────────────
    ws_ev = wb.create_sheet("Evidencias")
    ws_ev["A1"] = "EVIDENCIAS"
    ws_ev["A1"].font = Font(bold=True, size=12)
    ws_ev["A2"] = "Pantallazos de evidencia de ejecución de casos de prueba"
    ws_ev["A3"] = ""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Plantilla generada: {output_path}")


if __name__ == "__main__":
    # When running inside Docker container (/app is backend/)
    root = Path(__file__).parent.parent
    out = root / "templates" / "co-fr-vra-03" / "template.xlsx"
    make_template(out)
