"""Tests para app.deliverables.domain.excel_filler (RF-40)."""
import tempfile
from pathlib import Path

import pytest
import openpyxl

from app.deliverables.domain.excel_filler import fill_template


# ── Fixtures ──────────────────────────────────────────────────────────────────

def _make_template(tmp_path: Path) -> tuple[str, str]:
    """Crea plantilla y mapping mínimos para tests."""
    import yaml

    # Plantilla Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pruebas"
    ws["A1"] = "ORIGINAL_STYLE"
    ws["A1"].font = openpyxl.styles.Font(bold=True, size=14)
    wb.create_sheet("Evidencias")

    tmpl_path = str(tmp_path / "template.xlsx")
    wb.save(tmpl_path)

    # mapping.yaml
    mapping = {
        "fields": {
            "requirement_code": {"sheet": "Pruebas", "cell": "C5"},
            "module": {"sheet": "Pruebas", "cell": "C6"},
            "analyst": {"sheet": "Pruebas", "cell": "C7"},
            "date": {"sheet": "Pruebas", "cell": "C8"},
        },
        "tables": {
            "test_cases": {
                "sheet": "Pruebas",
                "start_row": 12,
                "columns": {
                    "code": "A",
                    "name": "B",
                    "result": "E",
                },
            }
        },
    }
    mapping_path = str(tmp_path / "mapping.yaml")
    with open(mapping_path, "w", encoding="utf-8") as f:
        yaml.dump(mapping, f)

    return tmpl_path, mapping_path


class TestFillTemplate:
    def test_fill_fields(self, tmp_path):
        tmpl, mapping = _make_template(tmp_path)
        out = str(tmp_path / "output.xlsx")
        fill_template(
            template_path=tmpl,
            mapping_path=mapping,
            data={"requirement_code": "IM-1234", "module": "Novedades", "analyst": "Ana"},
            output_path=out,
        )
        wb = openpyxl.load_workbook(out)
        ws = wb["Pruebas"]
        assert ws["C5"].value == "IM-1234"
        assert ws["C6"].value == "Novedades"
        assert ws["C7"].value == "Ana"

    def test_fill_table(self, tmp_path):
        tmpl, mapping = _make_template(tmp_path)
        out = str(tmp_path / "output.xlsx")
        cases = [
            {"code": "CP-01", "name": "Caso uno", "result": "passed"},
            {"code": "CP-02", "name": "Caso dos", "result": "failed"},
        ]
        fill_template(
            template_path=tmpl,
            mapping_path=mapping,
            data={"test_cases": cases},
            output_path=out,
        )
        wb = openpyxl.load_workbook(out)
        ws = wb["Pruebas"]
        assert ws["A12"].value == "CP-01"
        assert ws["B12"].value == "Caso uno"
        assert ws["A13"].value == "CP-02"
        assert ws["E13"].value == "failed"

    def test_preserves_styles(self, tmp_path):
        """El estilo de A1 (bold, size 14) no debe cambiar."""
        tmpl, mapping = _make_template(tmp_path)
        out = str(tmp_path / "output.xlsx")
        fill_template(
            template_path=tmpl,
            mapping_path=mapping,
            data={"requirement_code": "TEST"},
            output_path=out,
        )
        wb = openpyxl.load_workbook(out)
        ws = wb["Pruebas"]
        assert ws["A1"].font.bold is True
        assert ws["A1"].font.size == 14

    def test_missing_template_raises(self, tmp_path):
        _, mapping = _make_template(tmp_path)
        with pytest.raises(FileNotFoundError):
            fill_template(
                template_path="/no/existe.xlsx",
                mapping_path=mapping,
                data={},
                output_path=str(tmp_path / "out.xlsx"),
            )

    def test_steps_list_joined(self, tmp_path):
        tmpl, mapping = _make_template(tmp_path)
        out = str(tmp_path / "output.xlsx")
        cases = [{"code": "CP-01", "name": "Caso", "steps": ["Paso 1", "Paso 2"], "result": "passed"}]
        # Agregar columna steps al mapping
        import yaml
        with open(mapping, encoding="utf-8") as f:
            m = yaml.safe_load(f)
        m["tables"]["test_cases"]["columns"]["steps"] = "C"
        with open(mapping, "w", encoding="utf-8") as f:
            yaml.dump(m, f)

        fill_template(
            template_path=tmpl,
            mapping_path=mapping,
            data={"test_cases": cases},
            output_path=out,
        )
        wb = openpyxl.load_workbook(out)
        ws = wb["Pruebas"]
        assert "Paso 1" in (ws["C12"].value or "")
        assert "Paso 2" in (ws["C12"].value or "")
